"""
Обновляет final/BOM_Приложение_№1_(3).xlsx найденными ценами (30 позиций ИС).
Источник: Downloads/Приложение №1.xlsx (запрос от 2026-07-XX, Analog Devices/TI/Intel/Qorvo/Mini-Circuits)
Запускать: python update_bom6.py
"""
import urllib.request, json, re
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SCRIPT_DIR = Path(__file__).resolve().parent
OUTPUT = SCRIPT_DIR / "final" / "BOM_Приложение_№1_(3).xlsx"

def get_eur_usd():
    try:
        req = urllib.request.Request("https://open.er-api.com/v6/latest/EUR",
                                     headers={"User-Agent": "Mozilla/5.0"})
        with urllib.request.urlopen(req, timeout=10) as r:
            return float(json.loads(r.read().decode())["rates"]["USD"])
    except Exception:
        return 1.1362

def get_usd_rub(markup=4.0):
    try:
        req = urllib.request.Request("https://www.cbr.ru/scripts/XML_daily.asp",
                                     headers={"User-Agent": "Mozilla/5.0"})
        with urllib.request.urlopen(req, timeout=15) as r:
            xml = r.read().decode("cp1251", "replace")
        m = re.search(r'ID="R01235".*?<Nominal>(\d+)</Nominal>.*?<Value>([\d,]+)</Value>', xml, re.S)
        if m:
            return float(m.group(2).replace(",", ".")) / int(m.group(1)) + markup
    except Exception:
        pass
    return 79.6347 + markup

# ── Данные (собраны с oemsecrets через Chrome, 2026-07-02) ────────────────────

RESULTS = [
    {"pn": "AD9142ABCPZ", "qty": 3000, "manufacturer": "Analog Devices",
     "distributor": "Arrow", "moq": 1, "stock": 0, "lead": None,
     "price_eur": 62.2596, "status": "FOUND"},
    {"pn": "AD9513BCPZ", "qty": 600, "manufacturer": "Analog Devices",
     "distributor": "Arrow", "moq": 1, "stock": 36, "lead": None,
     "price_eur": 7.9065, "status": "FOUND"},
    {"pn": "AD9957BSVZ", "qty": 2300, "manufacturer": "Analog Devices",
     "distributor": "Mouser", "moq": 1, "stock": 0, "lead": None,
     "price_eur": 36.6000, "status": "FOUND"},
    {"pn": "ADL5375-05ACPZ-R7", "qty": 2900, "manufacturer": "Analog Devices",
     "distributor": "Arrow", "moq": 1, "stock": 2965, "lead": None,
     "price_eur": 7.1189, "status": "FOUND"},
    {"pn": "ADL5535ARKZ-R7", "qty": 9200, "manufacturer": "Analog Devices",
     "distributor": "Arrow", "moq": 1, "stock": 51, "lead": None,
     "price_eur": 2.3729, "status": "FOUND"},
    {"pn": "ADM7171ACPZ-1.8-R7", "qty": 1800, "manufacturer": "Analog Devices",
     "distributor": "Arrow", "moq": 1, "stock": 1089, "lead": None,
     "price_eur": 1.7923, "status": "FOUND"},
    {"pn": "ADM7171ACPZ-R7", "qty": 2900, "manufacturer": "Analog Devices",
     "distributor": "Arrow", "moq": 1, "stock": 34, "lead": None,
     "price_eur": 1.7923, "status": "FOUND"},
    {"pn": "ADP1707ARDZ-1.2-R7", "qty": 1200, "manufacturer": "Analog Devices",
     "distributor": "Arrow", "moq": 1, "stock": 0, "lead": None,
     "price_eur": 1.5892, "status": "FOUND"},
    {"pn": "ADP3335ACPZ-2.5-R7", "qty": 8100, "manufacturer": "Analog Devices",
     "distributor": "Arrow", "moq": 1, "stock": 0, "lead": None,
     "price_eur": 1.3545, "status": "FOUND"},
    {"pn": "ADP3335ACPZ-5-R7", "qty": 4100, "manufacturer": "Analog Devices",
     "distributor": "Arrow", "moq": 1, "stock": 0, "lead": None,
     "price_eur": 1.3261, "status": "FOUND"},
    {"pn": "ADP7156ACPZ-1.2-R7", "qty": 600, "manufacturer": "Analog Devices",
     "distributor": "Arrow", "moq": 1, "stock": 0, "lead": None,
     "price_eur": 4.0531, "status": "FOUND"},
    {"pn": "ADP7156ACPZ-1.8-R7", "qty": 2900, "manufacturer": "Analog Devices",
     "distributor": "Arrow", "moq": 1, "stock": 0, "lead": None,
     "price_eur": 4.0531, "status": "FOUND"},
    {"pn": "ADP7156ACPZ-3.3-R7", "qty": 1800, "manufacturer": "Analog Devices",
     "distributor": "Arrow", "moq": 1, "stock": 0, "lead": None,
     "price_eur": 4.0531, "status": "FOUND"},
    {"pn": "ADP7171ACPZ-1.8-R7", "qty": 600, "manufacturer": "Analog Devices", "status": "RFQ"},
    {"pn": "CBM3485AS8", "qty": 11000, "manufacturer": "Corebai", "status": "RFQ"},
    {"pn": "CDCLVC1106PW", "qty": 1200, "manufacturer": "Texas Instruments",
     "distributor": "Arrow", "moq": 1, "stock": 1327, "lead": None,
     "price_eur": 2.8467, "status": "FOUND"},
    {"pn": "EP3C16U4817N", "qty": 600, "manufacturer": "Intel", "status": "RFQ"},
    {"pn": "EP3C16U484I7N", "qty": 4100, "manufacturer": "Intel",
     "distributor": "Arrow", "moq": 1, "stock": 0, "lead": None,
     "price_eur": 121.9700, "status": "FOUND"},
    {"pn": "EPCS64SI16N", "qty": 4600, "manufacturer": "Intel",
     "distributor": "DigiKey", "moq": 1, "stock": 0, "lead": None,
     "price_eur": 32.7549, "status": "FOUND"},
    {"pn": "JPS-3-1W+", "qty": 1200, "manufacturer": "Mini-Circuits",
     "distributor": "Mouser", "moq": 1, "stock": 165, "lead": None,
     "price_eur": 20.7900, "status": "FOUND"},
    {"pn": "LMX2572RHAT", "qty": 2900, "manufacturer": "Texas Instruments",
     "distributor": "Arrow", "moq": 1, "stock": 0, "lead": None,
     "price_eur": 25.8690, "status": "FOUND"},
    {"pn": "LT1121IST-3.3", "qty": 2900, "manufacturer": "Analog Devices",
     "distributor": "Arrow", "moq": 1, "stock": 26, "lead": None,
     "price_eur": 3.5794, "status": "FOUND"},
    {"pn": "LT3481IMSE", "qty": 600, "manufacturer": "Analog Devices",
     "distributor": "Arrow", "moq": 1, "stock": 0, "lead": None,
     "price_eur": 5.3541, "status": "FOUND"},
    {"pn": "LT3504IUFD#TRPBF", "qty": 4700, "manufacturer": "Analog Devices",
     "distributor": "Mouser", "moq": 1, "stock": 0, "lead": None,
     "price_eur": 6.1400, "status": "FOUND"},
    {"pn": "LTC6957IMS-1#PBF", "qty": 2900, "manufacturer": "Analog Devices",
     "distributor": "Farnell", "moq": 1, "stock": 0, "lead": None,
     "price_eur": 7.1000, "status": "FOUND"},
    {"pn": "RF2052", "qty": 3500, "manufacturer": "Qorvo",
     "distributor": "Avnet America", "moq": 1, "stock": 0, "lead": None,
     "price_eur": 4.8219, "status": "FOUND"},
    {"pn": "RFDA4005SQ", "qty": 2900, "manufacturer": "Qorvo", "status": "RFQ"},
    {"pn": "SN74LVCC3245APWx", "qty": 2300, "manufacturer": "Texas Instruments",
     "distributor": "DigiKey", "moq": 1, "stock": 0, "lead": None,
     "bpn": "SN74LVCC3245APWR", "price_eur": 0.7138, "status": "FOUND"},
    {"pn": "SN74LVTH125IPWREP", "qty": 600, "manufacturer": "Texas Instruments",
     "distributor": "Arrow", "moq": 1, "stock": 2, "lead": None,
     "price_eur": 0.1611, "status": "FOUND"},
    {"pn": "TPS7A4700RGWx", "qty": 5800, "manufacturer": "Texas Instruments",
     "distributor": "Arrow", "moq": 1, "stock": 0, "lead": None,
     "bpn": "TPS7A4700RGWR", "price_eur": 2.6423, "status": "FOUND"},
]

# ── Стили ─────────────────────────────────────────────────────────────────────

F_HDR  = PatternFill("solid", fgColor="1F4E79")
F_HDRB = PatternFill("solid", fgColor="375623")
F_R1   = PatternFill("solid", fgColor="EBF3FB")
F_R2   = PatternFill("solid", fgColor="FFFFFF")
F_RFQ  = PatternFill("solid", fgColor="FFF2CC")
THIN   = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

HEADERS = [
    ("Part Number",              32, F_HDR),
    ("Distributor Part Number",  26, F_HDRB),
    ("Qty for Single BOM",       12, F_HDR),
    ("Manufacturer",             20, F_HDR),
    ("Distributor",              16, F_HDR),
    ("Minimum Order",            11, F_HDR),
    ("Stock",                    11, F_HDR),
    ("Lead Time (weeks)",        13, F_HDR),
    ("Unit Price USD",           14, F_HDR),
    ("Unit Price RUB",           15, F_HDR),
]

def write_cell(ws, row, col, value, fill, fmt=None, bold=False, color="000000"):
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(name="Arial", size=10, bold=bold, color=color)
    c.fill = fill
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = BORDER
    if fmt:
        c.number_format = fmt

def build_bom(rate_eur, rate_rub):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BOM"

    for col, (h, w, fill) in enumerate(HEADERS, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        c.fill = fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.row_dimensions[1].height = 32

    found_idx = 0
    for i, r in enumerate(RESULTS):
        row = i + 2
        ws.row_dimensions[row].height = 16
        pn = r["pn"]

        if r["status"] == "FOUND":
            base = F_R1 if found_idx % 2 == 0 else F_R2
            found_idx += 1
            stock = r.get("stock", 0)
            lead  = r.get("lead")
            in_stock = stock > 0 and lead is None

            if in_stock:
                lead_val = "In stock"
            elif lead is not None:
                lead_val = lead
            else:
                lead_val = ""

            if "price_usd" in r:
                usd = r["price_usd"]
            else:
                usd = round(r["price_eur"] * rate_eur, 5)
            rub = round(usd * rate_rub, 2)

            write_cell(ws, row, 1, pn,                   base)
            write_cell(ws, row, 2, r.get("bpn"),          base)
            write_cell(ws, row, 3, r["qty"],              base)
            write_cell(ws, row, 4, r.get("manufacturer"), base)
            write_cell(ws, row, 5, r.get("distributor"),  base)
            write_cell(ws, row, 6, r.get("moq"),          base)
            write_cell(ws, row, 7, stock,                 base)
            write_cell(ws, row, 8, lead_val,              base)
            write_cell(ws, row, 9, usd,                   base, "#,##0.0000")
            write_cell(ws, row,10, rub,                   base, "#,##0.00")
        else:
            for col in range(1, 11):
                v = None
                if col == 1: v = pn
                elif col == 3: v = r.get("qty")
                elif col == 4: v = r.get("manufacturer")
                elif col in (9, 10): v = "RFQ"
                bold = col in (9, 10)
                color = "7F4B00" if bold else "000000"
                write_cell(ws, row, col, v, F_RFQ, bold=bold, color=color)

    found = sum(1 for r in RESULTS if r["status"] == "FOUND")
    rfq   = len(RESULTS) - found
    note_row = len(RESULTS) + 3
    note = (
        f"Найдено: {found} | RFQ: {rfq} | Курс EUR/USD: {rate_eur:.4f} | "
        f"Курс USD/RUB (ЦБ + наценка): {rate_rub:.4f} | "
        f"Цена по ценовому брекету ≥ закупаемого кол-ва, минимум среди авторизованных. "
        f"Авторизованные: DigiKey, Mouser, Arrow, TTI, TME, Avnet, Newark/Farnell, Future. "
        f"Позиции 28, 30 (суффикс 'x'): найден реальный вариант упаковки 'R'. "
        f"Позиция 17 (EP3C16U4817N): вероятная опечатка PN (нет такого speed grade у Intel/Altera), "
        f"возможно имелась в виду EP3C16U484I7N (позиция 18) — уточнить у заказчика. "
        f"Позиции 14 (ADP7171ACPZ-1.8-R7), 27 (RFDA4005SQ): не найдены ни на одной площадке. "
        f"Позиция 15 (CBM3485AS8): китайский Corebai — нет у авторизованных дистрибьюторов."
    )
    nc = ws.cell(row=note_row, column=1, value=note)
    nc.font = Font(name="Arial", size=9, color="595959")
    nc.alignment = Alignment(horizontal="left", wrap_text=True)
    ws.merge_cells(f"A{note_row}:J{note_row}")
    ws.row_dimensions[note_row].height = 75

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)
    print(f"\nГотово: {OUTPUT}")
    print(f"Найдено: {found} | RFQ: {rfq}")
    print(f"EUR/USD: {rate_eur:.4f}  |  USD/RUB: {rate_rub:.4f}")

if __name__ == "__main__":
    print("Получаю актуальные курсы валют...")
    rate_eur = get_eur_usd()
    rate_rub = get_usd_rub()
    print(f"EUR/USD: {rate_eur:.4f}  |  USD/RUB: {rate_rub:.4f}")
    build_bom(rate_eur, rate_rub)
