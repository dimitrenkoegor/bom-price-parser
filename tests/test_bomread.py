# -*- coding: utf-8 -*-
"""Тесты чтения запроса: шапка, две колонки артикула/наименования, примечания, превью."""

import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

import openpyxl  # noqa: E402

from bomprice import bomread  # noqa: E402


def make_xlsx(path: Path, rows: list[list], sheet: str = "Лист1") -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    for row in rows:
        ws.append(row)
    wb.save(path)
    return path


class TestReadBom(unittest.TestCase):
    def test_combined_partnumber_column_and_brand_column(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = make_xlsx(Path(tmp) / "r.xlsx", [
                ["№", "Наименование", "Partnumber", "BRAND", "QTY"],
                [1, "Резонатор", "РЕЗОНАТОР Q 10,240-SS3-16-30/50-FU-T1 Jauch", "Jauch", 300],
            ])
            parts = bomread.read_bom(path, log=lambda *_: None)
        self.assertEqual(len(parts), 1)
        self.assertEqual(parts[0].mpn, "Q10,240-SS3-16-30/50-FU-T1")
        self.assertEqual(parts[0].manufacturer, "Jauch")
        self.assertEqual(parts[0].qty, 300)
        self.assertEqual(parts[0].raw, "Резонатор")

    def test_header_below_title_rows_and_notes_skipped(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = make_xlsx(Path(tmp) / "r.xlsx", [
                ["Запрос № 149", None, None, None],
                [None, None, None, None],
                ["№ п/п", "Условное обозначение", "Количество, шт", "Производитель"],
                [1, "Микросхема LM8272MM/NOPB Texas Instruments", 12, None],
                [None, "Примечание: срок поставки уточнить", None, None],
                [2, "GRM21BR71H105KA12", "1 000", "Murata"],
                [3, "Конденсатор", "по запросу", ""],
            ])
            parts = bomread.read_bom(path, log=lambda *_: None)
        self.assertEqual([p.mpn for p in parts], ["LM8272MM/NOPB", "GRM21BR71H105KA12"])
        self.assertEqual(parts[0].manufacturer, "Texas Instruments")
        self.assertEqual(parts[1].qty, 1000)
        self.assertEqual(parts[1].brand_source, "column")

    def test_brand_column_overrides_text(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = make_xlsx(Path(tmp) / "r.xlsx", [
                ["Part Number", "Qty", "Manufacturer"],
                ["BAS16H,115 NXP", 10, "Nexperia"],
            ])
            parts = bomread.read_bom(path, log=lambda *_: None)
        self.assertEqual(parts[0].manufacturer, "Nexperia")

    def test_missing_sheet_is_error_with_list(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = make_xlsx(Path(tmp) / "r.xlsx", [["Part Number", "Qty"], ["X1", 1]], sheet="BOM")
            with self.assertRaises(SystemExit) as ctx:
                bomread.read_bom(path, sheet="nope", log=lambda *_: None)
        self.assertIn("BOM", str(ctx.exception))

    def test_manual_columns_by_letter(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = make_xlsx(Path(tmp) / "r.xlsx", [
                ["a", "b", "c", "d"],
                ["x", "GRM21BR71H105KA12 Murata", "y", 5],
            ])
            parts = bomread.read_bom(path, header_row=1, pn_col=2, qty_col=4, log=lambda *_: None)
        self.assertEqual(parts[0].mpn, "GRM21BR71H105KA12")
        self.assertEqual(parts[0].qty, 5)


class TestPreviewRoundTrip(unittest.TestCase):
    def test_edits_survive(self):
        parts = [bomread.literal_part(1, "BAS16H", 100, "NXP", "Диод BAS16H"),
                 bomread.literal_part(2, "", 5, "", "Конденсатор непонятный")]
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "preview.xlsx"
            bomread.write_preview(parts, path)
            wb = openpyxl.load_workbook(path)
            ws = wb.active
            self.assertEqual(ws.cell(row=1, column=2).value, bomread.PREVIEW_COLS[1])
            ws.cell(row=3, column=2, value="GRM21BR71H105KA12")   # человек вписал артикул
            ws.cell(row=3, column=4, value="Murata")
            wb.save(path)
            back = bomread.read_preview(path)
        self.assertEqual(back[0].mpn, "BAS16H")
        self.assertEqual(back[0].manufacturer, "NXP Semiconductors")
        self.assertEqual(back[1].mpn, "GRM21BR71H105KA12")
        self.assertEqual(back[1].manufacturer, "Murata")
        self.assertEqual(back[1].qty, 5)
        self.assertTrue(back[1].variants)


class TestLiteralPart(unittest.TestCase):
    def test_once_mode_keeps_manufacturer_and_fixes_homoglyphs(self):
        part = bomread.literal_part(1, "ASSR-1530-005E", 200, "Avago")
        self.assertEqual(part.manufacturer, "Broadcom")
        self.assertEqual(part.qty, 200)
        part = bomread.literal_part(1, "М83513/01-AN", 1, "Amphenol")
        self.assertEqual(part.mpn, "M83513/01-AN")


if __name__ == "__main__":
    unittest.main(verbosity=2)
