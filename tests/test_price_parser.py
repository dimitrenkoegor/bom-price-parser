import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import openpyxl

import price_parser as parser


class PartNumberExtractionTests(unittest.TestCase):
    def test_jauch_part_keeps_prefix_decimal_and_slash(self):
        value = "РЕЗОНАТОР Q 10,240-SS3-16-30/50-FU-T1 Jauch"
        self.assertEqual(
            parser.extract_pn_from_description(value, "Jauch"),
            "Q 10,240-SS3-16-30/50-FU-T1",
        )

    def test_cyrillic_units_are_normalized_without_losing_model(self):
        self.assertEqual(
            parser.extract_pn_from_description("РЕЗОНАТОР МС-306 32,768КГЦ Epson", "Epson"),
            "MC-306 32,768kHz",
        )
        self.assertEqual(
            parser.extract_pn_from_description("РЕЗИСТОР 0603-200КОМ+-5% YAGEO", "YAGEO"),
            "0603-200K+-5%",
        )

    def test_explicit_catalog_part_wins_over_cyrillic_description(self):
        value = "РЕЗИСТОР 3310Y-001-10КОМ 3310Y001-103L Bourns"
        self.assertEqual(parser.extract_pn_from_description(value, "Bourns"), "3310Y001-103L")

    def test_decimal_comma_and_full_huber_suhner_part_are_preserved(self):
        self.assertEqual(
            parser.extract_pn_from_description("РЕЛЕ FTR-B4-CB-4,5Z FUJITSU", "FUJITSU"),
            "FTR-B4-CB-4,5Z",
        )
        self.assertEqual(
            parser.extract_pn_from_description("16 MMCX-50-2-13/111 OE Huber+Suhner", "Huber+Suhner"),
            "16 MMCX-50-2-13/111 OE",
        )

    def test_read_request_cleans_combined_partnumber_column(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "request.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append(["№", "Наименование", "Partnumber", "BRAND", "QTY"])
            ws.append([1, "Резонатор", "РЕЗОНАТОР Q 10,240-SS3-16-30/50-FU-T1 Jauch", "Jauch", 300])
            wb.save(path)

            items = parser.read_request_xlsx(path)

        self.assertEqual(items[0]["pn"], "Q 10,240-SS3-16-30/50-FU-T1")
        self.assertEqual(items[0]["manufacturer"], "Jauch")


class SearchValidationTests(unittest.TestCase):
    def test_mpn_match_allows_only_formatting_variation(self):
        self.assertTrue(parser.mpn_matches("Q 10,240-SS3", "Q10.240-SS3"))
        self.assertFalse(parser.mpn_matches("Q 10,240-SS3", "Q10.241-SS3"))

    def test_trailing_masks_per_regulament(self):
        # строчный x, смешанный Xx и заглавные XX — маска «любой суффикс»
        self.assertTrue(parser.mpn_matches("3313J-1-104x", "3313J-1-104E"))
        self.assertTrue(parser.mpn_matches("DLC70B8R2CW501Xx", "DLC70B8R2CW501XT"))
        self.assertTrue(parser.mpn_matches("DLC70B8R2CW501XX", "DLC70B8R2CW501XK"))
        # одиночная заглавная X — литерал, не маска: произвольный (неупаковочный)
        # суффикс не принимается; T принялся бы по правилу упаковки — это норм
        self.assertFalse(parser.mpn_matches("EXB28V220JX", "EXB28V220JXDR"))
        self.assertTrue(parser.mpn_matches("EXB28V220JX", "EXB-28V220JX"))

    def test_midstring_x_wildcard(self):
        self.assertTrue(parser.mpn_matches("TPSE477K010x0200", "TPSE477K010R0200"))
        self.assertFalse(parser.mpn_matches("TPSE477K010x0200", "TPSE477K0100200"))
        self.assertFalse(parser.mpn_matches("TPSE477K010x0200", "TPSE477K010RR0200"))

    def test_packaging_tail_and_series_prefix(self):
        self.assertTrue(parser.mpn_matches("CR0805-JW-390E", "CR0805-JW-390ELF"))
        self.assertTrue(parser.mpn_matches("0805-FX-1503", "CR0805-FX-1503ELF"))
        self.assertFalse(parser.mpn_matches("LM358", "LM358DR"))
        self.assertFalse(parser.mpn_matches("LM358", "PLM358"))
        self.assertFalse(parser.mpn_matches("104", "CR104"))

    def test_manufacturer_aliases_cover_merged_brands(self):
        self.assertTrue(parser.manufacturer_matches("Epcos", "TDK Electronics / EPCOS"))
        self.assertTrue(parser.manufacturer_matches("Avago Technologies", "Broadcom Limited"))
        self.assertFalse(parser.manufacturer_matches("Jauch", "Yageo"))

    def test_price_search_rejects_other_manufacturer(self):
        wrong = {
            "dist": "DigiKey", "pn": "WRONG-1", "stock": 1000, "moq": 1,
            "lead": None, "tiers": [{"qty": 1, "price": 0.01}], "manu": "Yageo",
        }
        correct = {
            "dist": "DigiKey", "pn": "RIGHT-1", "stock": 1000, "moq": 1,
            "lead": None, "tiers": [{"qty": 1, "price": 0.02}], "manu": "Jauch",
        }

        def digikey(query, _qty, match_pn=None):
            return [correct] if "." in query else [wrong]

        with patch.object(parser, "search_digikey", side_effect=digikey), \
                patch.object(parser, "search_mouser", return_value=[]), \
                patch.object(parser, "search_farnell", return_value=[]), \
                patch.object(parser, "search_tme", return_value=[]):
            result = parser.find_best_price("Q 10,240-SS3", 300, "Jauch")

        self.assertEqual(result["status"], "FOUND")
        self.assertEqual(result["manufacturer"], "Jauch")
        self.assertEqual(result["distr_pn"], "RIGHT-1")

    def test_transient_timeout_is_retried(self):
        class Response:
            def __enter__(self):
                return self

            def __exit__(self, *_args):
                return False

            def read(self):
                return b"ok"

        with patch("urllib.request.urlopen", side_effect=[TimeoutError(), Response()]), \
                patch.object(parser.time, "sleep"):
            self.assertEqual(parser.fetch_bytes("request"), b"ok")

    def test_once_mode_keeps_required_manufacturer(self):
        args = SimpleNamespace(once=["ASSR-1530-005E", "200"], manufacturer="Avago")
        self.assertEqual(parser.load_items(args)[0]["manufacturer"], "Avago")


class OutputFootnoteTests(unittest.TestCase):
    def test_llm_row_gets_peach_fill_and_is_counted_in_footnote(self):
        results = [
            {"pn": "LM358D", "qty": 10, "manufacturer": "TI", "description": "LM358D",
             "status": "FOUND", "distributor": "Mouser", "price_usd": 0.5,
             "stock": 100, "moq": 1, "lead": None, "in_stock": True,
             "resolved_mpn": "", "distr_pn": ""},
            {"pn": "MAX481ESA", "qty": 2, "manufacturer": "Maxim", "description": "x",
             "status": "FOUND", "distributor": "DigiKey", "price_usd": 6.74,
             "stock": 10, "moq": 1, "lead": None, "in_stock": True,
             "resolved_mpn": "MAX481ESA+T", "distr_pn": "", "llm_assisted": True},
            {"pn": "MYSTERY-1", "qty": 5, "manufacturer": "",
             "description": "MYSTERY-1", "status": "RFQ"},
        ]
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "out.xlsx"
            parser.write_results(results, path, 80.0)
            ws = openpyxl.load_workbook(path).active
            self.assertEqual(str(ws.cell(row=3, column=1).fill.fgColor.rgb)[-6:], "FDE9D9")
            self.assertEqual(ws.cell(row=3, column=2).value, "MAX481ESA+T")
            note = ws.cell(row=len(results) + 3, column=1).value
        self.assertIn("Найдено: 2 | RFQ: 1", note)
        self.assertIn("только официальные API", note)
        self.assertIn("подтверждено у дистрибьютора: 1", note)


if __name__ == "__main__":
    unittest.main()
