# -*- coding: utf-8 -*-
"""Тесты итогового xlsx: 10 колонок, заливки, колонка B, сноска, строки для монитора."""

import io
import sys
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

import openpyxl  # noqa: E402

from bomprice import fx as fxmod, report  # noqa: E402

RATES = fxmod.Rates(usd_rub=84.0, usd_rub_cbr=80.0, markup=4.0, per_usd={}, at="2026-09-09")


def found(**kw):
    base = {"status": "FOUND", "pn": "BAS16H", "qty": 100, "manufacturer": "Nexperia",
            "distributor": "DigiKey", "distr_pn": "1727-BAS16HCT-ND", "resolved_mpn": "BAS16H,115",
            "moq": 1, "stock": 5000, "in_stock": True, "lead": None, "price_usd": 0.123,
            "llm_assisted": False}
    base.update(kw)
    return base


RFQ = {"status": "RFQ", "pn": "XYZ-1", "qty": 3, "manufacturer": "", "description": "XYZ-1"}


class TestWorkbook(unittest.TestCase):
    def _build(self, results):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "out.xlsx"
            buf = io.StringIO()
            with redirect_stdout(buf):
                report.write_results(results, path, RATES, ["DigiKey", "Mouser"])
            wb = openpyxl.load_workbook(path)
            return wb.active, buf.getvalue()

    def test_columns_and_monitor_lines(self):
        ws, out = self._build([found(), RFQ])
        self.assertEqual([c.value for c in ws[1]][:10], [h for h, _ in report.HEADERS])
        self.assertEqual(ws.cell(row=2, column=2).value, "BAS16H,115")       # колонка B: суффикс
        self.assertEqual(ws.cell(row=2, column=8).value, "In stock")
        self.assertAlmostEqual(ws.cell(row=2, column=10).value, round(0.123 * 84.0, 2))
        self.assertEqual(ws.cell(row=3, column=9).value, "RFQ")
        self.assertEqual(ws.cell(row=3, column=1).fill.fgColor.rgb[-6:], report.FILL_RFQ)
        self.assertIn("Найдено: 1 | RFQ: 1", out)
        self.assertIn("Готово:", out)

    def test_column_b_empty_when_same_pn(self):
        ws, _ = self._build([found(resolved_mpn="BAS-16H")])
        self.assertIsNone(ws.cell(row=2, column=2).value)

    def test_lead_text_variants(self):
        self.assertEqual(report.lead_text(found(in_stock=False, lead=1, stock=10)), "1 week")
        self.assertEqual(report.lead_text(found(in_stock=False, lead=6, stock=10)), "6 weeks")
        self.assertEqual(report.lead_text(found(in_stock=False, lead=None, stock=0)), "уточнить у поставщика")

    def test_footnote_mentions_moq_and_llm(self):
        text = report.footnote([found(moq=500), found(pn="A1", llm_assisted=True), RFQ], RATES, ["DigiKey"])
        self.assertIn("Найдено: 2 | RFQ: 1", text)
        self.assertIn("MOQ выше запрошенного", text)
        self.assertIn("BAS16H (MOQ 500)", text)
        self.assertIn("персиковым", text)
        self.assertIn("84.0000", text)

    def test_llm_rows_are_peach(self):
        ws, _ = self._build([found(llm_assisted=True)])
        self.assertEqual(ws.cell(row=2, column=1).fill.fgColor.rgb[-6:], report.FILL_LLM)


if __name__ == "__main__":
    unittest.main(verbosity=2)
