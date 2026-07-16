"""
Автономный монитор почты для price_parser.py — без Mission Control и моста.

Цикл: проверяет ящик по IMAP → письмо от разрешённого отправителя с командой
(«найти цены», «просчитать BOM», «дай цены», «прайс») и вложением Excel/TXT →
прогоняет вложение через price_parser.py → отвечает на письмо готовым файлом
(SEND_ENABLED=true) → RFQ-остаток пишет в jobs_mail/<job>/rfq_list.txt
(его потом можно добить вручную через Claude + oemsecrets).

Настройки берутся из .env рядом со скриптом (дополнительно к ключам API):
    MAIL_IMAP_HOST=imap.mail.ru
    MAIL_USER=e.dimitrenko@chainchip.ru
    MAIL_PASSWORD=<app-пароль>
    MAIL_SMTP_HOST=smtp.mail.ru
    MAIL_SMTP_PORT=465
    ALLOWED_SENDERS=e.dimitrenko@chainchip.ru        # через ; можно @домен
    POLL_SECONDS=60
    SEND_ENABLED=true
    ENABLE_CLAUDE_FALLBACK=false   # true после `claude` /login — советы по RFQ

Запуск:
    python mail_monitor.py            # бесконечный цикл (демон)
    python mail_monitor.py --once     # один цикл (проверка)
    python mail_monitor.py --once --dry-run   # без отправки писем

Остановка демона: файл STOP_MONITOR рядом со скриптом (создаёт STOP_MAIL.bat)
или Ctrl+C. ВАЖНО: не запускайте одновременно с agent-mail-bridge — оба
слушают один ящик и будут отбирать письма друг у друга.
"""

from __future__ import annotations

import argparse
import email
import hashlib
import imaplib
import json
import mimetypes
import re
import smtplib
import ssl
import subprocess
import sys
import time
from datetime import datetime
from email.header import decode_header, make_header
from email.message import EmailMessage, Message
from email.utils import formataddr, parseaddr
from pathlib import Path

SCRIPT_DIR = Path(__file__).resolve().parent
STATE_PATH = SCRIPT_DIR / "mail-state.json"
STOP_FLAG = SCRIPT_DIR / "STOP_MONITOR"
JOB_ROOT = SCRIPT_DIR / "jobs_mail"
PARSER = SCRIPT_DIR / "price_parser.py"
SUPPORTED = {".xlsx", ".xlsm", ".xls", ".txt"}
TRIGGER = re.compile(r"(?:найти|дай|собери|собрать)\s*цены|просчитать\s*bom|прайс|цены\s*с\s*площад", re.IGNORECASE)


def load_env() -> dict[str, str]:
    env: dict[str, str] = {}
    path = SCRIPT_DIR / ".env"
    if path.exists():
        for raw in path.read_text(encoding="utf-8-sig").splitlines():
            line = raw.strip()
            if line and not line.startswith("#") and "=" in line:
                key, _, value = line.partition("=")
                env[key.strip()] = value.strip()
    return env


ENV = load_env()


def env_bool(name: str, default: bool = False) -> bool:
    return ENV.get(name, str(default)).strip().lower() in {"1", "true", "yes", "on"}


def log(message: str) -> None:
    line = f"{datetime.now():%Y-%m-%d %H:%M:%S} {message}"
    print(line, flush=True)
    with open(SCRIPT_DIR / "monitor.log", "a", encoding="utf-8") as handle:
        handle.write(line + "\n")


def load_state() -> dict:
    try:
        return json.loads(STATE_PATH.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return {"processed": [], "inspected": []}


def save_state(state: dict) -> None:
    state["processed"] = state["processed"][-1000:]
    state["inspected"] = state["inspected"][-2000:]
    tmp = STATE_PATH.with_suffix(".tmp")
    tmp.write_text(json.dumps(state, ensure_ascii=False, indent=1), encoding="utf-8")
    tmp.replace(STATE_PATH)


def decode(value) -> str:
    return str(make_header(decode_header(value or "")))


def sender_allowed(sender: str) -> bool:
    sender = sender.lower()
    for entry in ENV.get("ALLOWED_SENDERS", "").split(";"):
        entry = entry.strip().lower()
        if not entry:
            continue
        if entry.startswith("@") and sender.endswith(entry):
            return True
        if sender == entry:
            return True
    return False


def body_text(message: Message) -> str:
    if message.is_multipart():
        for part in message.walk():
            disposition = str(part.get("Content-Disposition", "")).lower()
            if part.get_content_type() == "text/plain" and "attachment" not in disposition:
                payload = part.get_payload(decode=True) or b""
                return payload.decode(part.get_content_charset() or "utf-8", errors="replace")
        return ""
    payload = message.get_payload(decode=True) or b""
    return payload.decode(message.get_content_charset() or "utf-8", errors="replace")


def safe_name(name: str, fallback: str) -> str:
    name = Path(name).name
    return re.sub(r"[^A-Za-zА-Яа-яЁё0-9._() -]", "_", name)[:150] or fallback


def save_attachments(message: Message, destination: Path) -> list[Path]:
    destination.mkdir(parents=True, exist_ok=True)
    saved: list[Path] = []
    for index, part in enumerate(message.walk(), 1):
        filename = part.get_filename()
        if not filename:
            continue
        filename = decode(filename)
        if Path(filename).suffix.lower() not in SUPPORTED:
            continue
        target = destination / safe_name(filename, f"attachment-{index}")
        counter = 1
        while target.exists():
            target = destination / f"{target.stem}-{counter}{target.suffix}"
            counter += 1
        target.write_bytes(part.get_payload(decode=True) or b"")
        saved.append(target)
    return saved


def run_parser(source: Path, job_dir: Path) -> tuple[Path | None, str]:
    """Возвращает (готовый файл | None, строка-итог для письма)."""
    stem = re.sub(r"[^A-Za-zА-Яа-яЁё0-9_-]", "_", source.stem)[:60]
    output = job_dir / f"BOM_{stem}.xlsx"
    command = [
        sys.executable, str(PARSER), "--input", str(source), "--output", str(output),
        "--preview-path", str(job_dir / f"preview_{stem}.xlsx"), "--yes",
    ]
    import os
    child_env = {**os.environ, "PYTHONIOENCODING": "utf-8"}  # иначе stdout в cp1251 и итог не парсится
    completed = subprocess.run(command, cwd=SCRIPT_DIR, text=True, env=child_env,
                               capture_output=True, timeout=1800, encoding="utf-8", errors="replace")
    (job_dir / f"parser_{stem}.log").write_text(
        (completed.stdout or "") + "\n" + (completed.stderr or ""), encoding="utf-8")
    if completed.returncode or not output.exists():
        return None, f"{source.name}: ошибка обработки (см. лог на сервере)"
    match = re.search(r"Найдено:\s*(\d+)\s*\|\s*RFQ:\s*(\d+)", completed.stdout or "")
    summary = f"{source.name}: найдено {match.group(1)}, RFQ {match.group(2)}" if match else f"{source.name}: обработан"
    return output, summary


def rfq_parts_from_workbook(path: Path) -> list[str]:
    import openpyxl

    worksheet = openpyxl.load_workbook(path, data_only=True).active
    parts: list[str] = []
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        if row and row[0] and len(row) >= 9 and row[2] and not row[4]:
            parts.append(str(row[0]).strip())
    return parts


def sanity_check_workbook(path: Path) -> str:
    """Проверка результата перед автоотправкой. Возвращает '' если всё в порядке,
    иначе описание проблемы (письмо клиенту в этом случае НЕ уходит)."""
    import openpyxl

    try:
        worksheet = openpyxl.load_workbook(path, data_only=True).active
    except Exception as exc:
        return f"файл результата не читается: {exc}"
    rows = found = 0
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        if not row or not row[0] or len(row) < 10 or not row[2]:
            continue
        rows += 1
        if row[4]:  # найденная позиция: цена обязана быть числом > 0
            found += 1
            price = row[8]
            if not isinstance(price, (int, float)) or price <= 0:
                return f"позиция {row[0]}: некорректная цена {price!r}"
    if rows == 0:
        return "в результате нет ни одной позиции"
    return ""


def notify_self(subject: str, text: str) -> None:
    """Служебное письмо на собственный ящик при сбое (не клиенту)."""
    try:
        user = ENV["MAIL_USER"]
        message = EmailMessage()
        message["From"] = f"BOM Price Monitor <{user}>"
        message["To"] = user
        message["Subject"] = f"[MONITOR] {subject}"
        message.set_content(text)
        host = ENV.get("MAIL_SMTP_HOST", "smtp.mail.ru")
        port = int(ENV.get("MAIL_SMTP_PORT", "465"))
        with smtplib.SMTP_SSL(host, port, context=ssl.create_default_context()) as smtp:
            smtp.login(user, ENV["MAIL_PASSWORD"])
            smtp.send_message(message)
    except Exception as exc:
        log(f"notify_self не удалось: {exc}")


def claude_advisory(job_dir: Path, rfq_parts: list[str]) -> None:
    """Необязательный советник: headless Claude пишет JSON-подсказки по RFQ."""
    if not env_bool("ENABLE_CLAUDE_FALLBACK"):
        return
    prompt = (
        "Ты — fallback-исследователь цен на электронные компоненты. Для каждого PN "
        "из списка предложи, где искать среди АВТОРИЗОВАННЫХ дистрибьюторов "
        "(без брокеров), и заметь вероятные ошибки записи артикула. Верни строгий "
        "JSON-массив объектов {pn, note, suggested_pn, suggested_sources}. Список: "
        + json.dumps(rfq_parts[:100], ensure_ascii=False)
    )
    try:
        completed = subprocess.run(["claude", "-p", prompt, "--output-format", "json"],
                                   cwd=job_dir, text=True, capture_output=True,
                                   timeout=300, encoding="utf-8", errors="replace")
        (job_dir / "fallback.json").write_text(completed.stdout or completed.stderr or "", encoding="utf-8")
    except (OSError, subprocess.TimeoutExpired) as exc:
        (job_dir / "fallback.json").write_text(f"Fallback unavailable: {exc}", encoding="utf-8")


def send_reply(original: Message, sender: str, results: list[Path], summaries: list[str], rfq_total: int) -> None:
    user = ENV["MAIL_USER"]
    reply = EmailMessage()
    reply["From"] = formataddr(("BOM Price Monitor", user))
    reply["To"] = sender
    subject = decode(original.get("Subject")) or "запрос цен"
    reply["Subject"] = f"Re: {subject}"
    message_id = original.get("Message-ID")
    if message_id:
        reply["In-Reply-To"] = message_id
        reply["References"] = message_id
    lines = ["Здравствуйте!", "", "Автоматический расчёт цен выполнен, файлы во вложении:", ""]
    lines += [f"  - {s}" for s in summaries]
    if rfq_total:
        lines += ["", f"Позиции RFQ ({rfq_total} шт.) отмечены в файле жёлтым — по ним нужен ручной запрос поставщику."]
    lines += ["", "--", "BOM Price Monitor (автоматический ответ)"]
    reply.set_content("\n".join(lines))
    for path in results:
        ctype, _ = mimetypes.guess_type(path.name)
        maintype, subtype = (ctype or "application/octet-stream").split("/", 1)
        reply.add_attachment(path.read_bytes(), maintype=maintype, subtype=subtype, filename=path.name)
    host = ENV.get("MAIL_SMTP_HOST", "smtp.mail.ru")
    port = int(ENV.get("MAIL_SMTP_PORT", "465"))
    with smtplib.SMTP_SSL(host, port, context=ssl.create_default_context()) as smtp:
        smtp.login(user, ENV["MAIL_PASSWORD"])
        smtp.send_message(reply)


def intake_once(dry_run: bool = False) -> None:
    state = load_state()
    processed = set(state["processed"])
    inspected = set(state["inspected"])
    host = ENV.get("MAIL_IMAP_HOST", "imap.mail.ru")
    with imaplib.IMAP4_SSL(host, 993, ssl_context=ssl.create_default_context()) as client:
        client.login(ENV["MAIL_USER"], ENV["MAIL_PASSWORD"])
        client.select(ENV.get("MAIL_IMAP_MAILBOX", "INBOX"))
        status, rows = client.uid("search", None, "UNSEEN")
        if status != "OK":
            raise RuntimeError("IMAP search failed")
        for uid in rows[0].split()[-50:]:
            status, data = client.uid("fetch", uid, "(BODY.PEEK[HEADER.FIELDS (FROM SUBJECT MESSAGE-ID)])")
            if status != "OK" or not data or data[0] is None:
                continue
            header = email.message_from_bytes(data[0][1])
            sender = parseaddr(header.get("From", ""))[1].lower()
            subject = decode(header.get("Subject"))
            message_id = header.get("Message-ID") or f"uid:{uid.decode()}:{sender}:{subject}"
            fingerprint = hashlib.sha256(message_id.encode()).hexdigest()
            if fingerprint in inspected or not sender_allowed(sender):
                continue
            status, data = client.uid("fetch", uid, "(BODY.PEEK[])")
            if status != "OK" or not data or data[0] is None:
                continue
            message = email.message_from_bytes(data[0][1])
            if not TRIGGER.search(f"{subject}\n{body_text(message)}"):
                inspected.add(fingerprint)
                continue

            job_dir = JOB_ROOT / f"{datetime.now():%Y%m%d-%H%M%S}-{fingerprint[:8]}"
            attachments = save_attachments(message, job_dir / "source")
            log(f"Письмо от {sender}: '{subject[:80]}' — вложений: {len(attachments)}")
            results: list[Path] = []
            summaries: list[str] = []
            problems: list[str] = []
            rfq_total = 0
            for attachment in attachments:
                output, summary = run_parser(attachment, job_dir)
                summaries.append(summary)
                if output:
                    issue = sanity_check_workbook(output)
                    if issue:
                        problems.append(f"{attachment.name}: {issue}")
                        continue                       # файл с проблемой клиенту не уходит
                    results.append(output)
                    try:
                        rfq_parts = rfq_parts_from_workbook(output)
                    except Exception:
                        rfq_parts = []
                    if rfq_parts:
                        rfq_total += len(rfq_parts)
                        with open(job_dir / "rfq_list.txt", "a", encoding="utf-8") as handle:
                            handle.write(f"# {attachment.name}\n" + "\n".join(rfq_parts) + "\n")
                else:
                    problems.append(summary)
            if not attachments:
                summaries.append("Не найдено вложение Excel/TXT — пришлите BOM файлом.")
            if rfq_total:
                claude_advisory(job_dir, (job_dir / "rfq_list.txt").read_text(encoding="utf-8").splitlines())

            if problems:
                log(f"САНИТИ-ПРОБЛЕМЫ (не отправлено клиенту): {'; '.join(problems)}")
                notify_self("сбой обработки задания",
                            f"Письмо от {sender}: '{subject}'\nПапка: {job_dir}\n\nПроблемы:\n- "
                            + "\n- ".join(problems))

            send_enabled = env_bool("SEND_ENABLED", True) and not dry_run
            if send_enabled and results:
                send_reply(message, sender, results, summaries, rfq_total)
                log(f"Ответ отправлен -> {sender}: {'; '.join(summaries)}")
            elif send_enabled and not attachments:
                send_reply(message, sender, [], summaries, 0)
                log(f"Ответ (нет вложений) -> {sender}")
            else:
                log(f"[dry-run/off/сбой] Ответ НЕ отправлен. Итог: {'; '.join(summaries)}. Файлы: {job_dir}")

            processed.add(fingerprint)
            inspected.add(fingerprint)
            state["processed"] = sorted(processed)
            state["inspected"] = sorted(inspected)
            save_state(state)
            client.uid("store", uid, "+FLAGS", "(\\Seen)")
    state["processed"] = sorted(processed)
    state["inspected"] = sorted(inspected)
    save_state(state)


def main() -> int:
    parser = argparse.ArgumentParser(description="Монитор почты для price_parser")
    parser.add_argument("--once", action="store_true", help="один цикл и выход")
    parser.add_argument("--dry-run", action="store_true", help="не отправлять письма")
    args = parser.parse_args()

    for required in ("MAIL_USER", "MAIL_PASSWORD"):
        if not ENV.get(required):
            print(f"В .env отсутствует {required}")
            return 1
    if not ENV.get("ALLOWED_SENDERS"):
        print("В .env отсутствует ALLOWED_SENDERS (белый список отправителей)")
        return 1

    JOB_ROOT.mkdir(exist_ok=True)
    poll = max(30, int(ENV.get("POLL_SECONDS", "60")))
    log(f"Монитор запущен (poll={poll}s, send={'ON' if env_bool('SEND_ENABLED', True) and not args.dry_run else 'OFF'})")
    while True:
        if STOP_FLAG.exists():
            log("Обнаружен STOP_MONITOR — останавливаюсь.")
            STOP_FLAG.unlink(missing_ok=True)
            return 0
        try:
            intake_once(dry_run=args.dry_run)
        except Exception as exc:
            log(f"Цикл завершился ошибкой (повтор через {poll}s): {exc}")
        if args.once:
            return 0
        time.sleep(poll)


if __name__ == "__main__":
    raise SystemExit(main())
