"""Итоговый xlsx: 10 колонок по эталону заказчика + сноска. Контракт — .agents/rules/output-format.md.

Печатает две строки, которые парсит почтовый монитор: «Готово: <путь>» и
«Найдено: N | RFQ: M». Менять их формат нельзя.
"""

from __future__ import annotations

from pathlib import Path

from . import fx as fxmod
from . import mpn as mpnmod

HEADERS = [
    ("Part Number", 30), ("Distributor Part Number", 26), ("Qty for Single BOM", 11),
    ("Manufacturer", 20), ("Distributor", 20), ("Minimum Order", 11), ("Stock", 10),
    ("Lead Time (weeks)", 13), ("Unit Price USD", 14), ("Unit Price RUB", 15),
]

FILL_HEAD = "1F4E79"
FILL_HEAD_B = "375623"
FILL_ROW_A = "EBF3FB"
FILL_ROW_B = "FFFFFF"
FILL_RFQ = "FFF2CC"
FILL_COL_B = "E2EFDA"
FILL_LLM = "FDE9D9"
COLOR_RFQ = "7F4B00"


def column_b(result: dict) -> str:
    """Колонка B — конкретный артикул, когда он отличается от запрошенного
    (маска, суффикс упаковки, префикс серии, написание)."""
    pn = result.get("pn") or ""
    resolved = result.get("resolved_mpn") or ""
    if resolved and mpnmod.canonical(resolved) != mpnmod.canonical(pn):
        return resolved
    return ""


def lead_text(result: dict) -> str:
    lead = result.get("lead")
    if result.get("in_stock") and not lead:
        return "In stock"
    if lead is not None:
        return f"{lead} weeks" if lead != 1 else "1 week"
    if not result.get("stock"):
        return "уточнить у поставщика"
    return ""


def footnote(results: list[dict], rates: fxmod.Rates, source_titles: list[str]) -> str:
    found = [r for r in results if r.get("status") == "FOUND"]
    rfq = len(results) - len(found)
    text = (f"Найдено: {len(found)} | RFQ: {rfq} | "
            f"Курс USD/RUB (ЦБ + наценка): {rates.usd_rub:.4f} | "
            f"Цена по ценовому брекету ≥ количества, приоритет позиций в наличии. "
            f"Источник цен — только официальные API дистрибьюторов "
            f"({', '.join(source_titles)}).")
    moq_rows = [r for r in found if int(r.get("moq") or 1) > int(r.get("qty") or 0)]
    if moq_rows:
        text += (" MOQ выше запрошенного количества у позиций: "
                 + ", ".join(f"{r['pn']} (MOQ {r['moq']})" for r in moq_rows[:20])
                 + (" …" if len(moq_rows) > 20 else "")
                 + " — закупать придётся минимальную партию.")
    llm_rows = sum(1 for r in found if r.get("llm_assisted"))
    if llm_rows:
        text += (f" Позиций, где написание артикула уточнено автоматически и затем "
                 f"подтверждено у дистрибьютора: {llm_rows} — выделены персиковым, "
                 f"рекомендуется выборочная проверка.")
    return text


def write_results(results: list[dict], output_path: Path, rates: fxmod.Rates,
                  source_titles: list[str], log=print) -> Path:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BOM"

    fills = {name: PatternFill("solid", fgColor=color) for name, color in (
        ("head", FILL_HEAD), ("head_b", FILL_HEAD_B), ("a", FILL_ROW_A), ("b", FILL_ROW_B),
        ("rfq", FILL_RFQ), ("col_b", FILL_COL_B), ("llm", FILL_LLM))}
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col, (title, width) in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = fills["head_b"] if col == 2 else fills["head"]
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[1].height = 32

    found_idx = 0
    for i, r in enumerate(results):
        row = i + 2
        ws.row_dimensions[row].height = 16
        pn = r.get("pn") or r.get("description") or ""
        if r.get("status") == "FOUND":
            base = fills["llm"] if r.get("llm_assisted") else (
                fills["a"] if found_idx % 2 == 0 else fills["b"])
            found_idx += 1
            b_val = column_b(r)
            usd = r.get("price_usd")
            rub = rates.to_rub(usd, "USD") if usd is not None else None
            vals = [
                (pn, base, None, False, "000000"),
                (b_val or None, fills["col_b"] if b_val else base, None, False, "000000"),
                (r.get("qty"), base, None, False, "000000"),
                (r.get("manufacturer") or None, base, None, False, "000000"),
                (r.get("distributor"), base, None, False, "000000"),
                (r.get("moq"), base, None, False, "000000"),
                (r.get("stock"), base, None, False, "000000"),
                (lead_text(r), base, None, False, "000000"),
                (usd, base, "#,##0.000", False, "000000"),
                (rub, base, "#,##0.00", False, "000000"),
            ]
        else:
            f = fills["rfq"]
            vals = [
                (pn, f, None, False, "000000"), (None, f, None, False, "000000"),
                (r.get("qty"), f, None, False, "000000"),
                (r.get("manufacturer") or None, f, None, False, "000000"),
                (None, f, None, False, "000000"), (None, f, None, False, "000000"),
                (None, f, None, False, "000000"), (None, f, None, False, "000000"),
                ("RFQ", f, None, True, COLOR_RFQ), ("RFQ", f, None, True, COLOR_RFQ),
            ]
        for col, (value, fill, fmt, bold, color) in enumerate(vals, 1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.font = Font(name="Arial", size=10, bold=bold, color=color)
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border
            if fmt:
                cell.number_format = fmt

    note_row = len(results) + 3
    cell = ws.cell(row=note_row, column=1, value=footnote(results, rates, source_titles))
    cell.font = Font(name="Arial", size=9, color="595959")
    cell.alignment = Alignment(horizontal="left", wrap_text=True)
    ws.merge_cells(f"A{note_row}:J{note_row}")
    ws.row_dimensions[note_row].height = 45

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)

    found = sum(1 for r in results if r.get("status") == "FOUND")
    log(f"\nГотово: {output_path}")
    log(f"Найдено: {found} | RFQ: {len(results) - found}")
    return output_path
