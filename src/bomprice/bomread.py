"""Чтение BOM-запроса (xlsx / txt) → позиции, превью для человека и обратно.

Шапка и колонки ищутся автоматически. В запросах бывают две колонки — «Условное
обозначение» (артикул) и «Наименование» (описание); артикул берётся из первой,
если она есть, бренд — из колонки «Производитель», иначе из текста.
Строки без количества (примечания, подзаголовки) пропускаются.
"""

from __future__ import annotations

import re
from pathlib import Path

from . import mpn as mpnmod
from . import store

PN_HINTS = ("условное обознач", "part number", "partnumber", "артикул", "парт",
            "p/n", "mpn", "обозначение")
NAME_HINTS = ("наименован", "товар", "описан", "part name", "description", "позиция",
              "номенклатура")
QTY_HINTS = ("количеств", "кол-во", "колво", "qty", "q-ty", "q ty", "quantity")
MFR_HINTS = ("производител", "изготовит", "manufacturer", "mfr", "brand", "бренд")
NUM_HINTS = ("№", "n п/п", "п/п")


def _norm(value) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip().lower()


def _find_header(ws, scan_rows: int = 40) -> dict:
    """→ {row, pn, name, qty, mfr, num} — номера колонок (0 = нет)."""
    fallback = None
    for row_idx, row in enumerate(
            ws.iter_rows(min_row=1, max_row=min(scan_rows, ws.max_row)), 1):
        cols = {"pn": 0, "name": 0, "qty": 0, "mfr": 0, "num": 0}
        for col_idx, cell in enumerate(row, 1):
            text = _norm(getattr(cell, "value", None))
            if not text:
                continue
            if not cols["num"] and any(h in text for h in NUM_HINTS) and len(text) <= 6:
                cols["num"] = col_idx
            elif not cols["pn"] and any(h in text for h in PN_HINTS):
                cols["pn"] = col_idx
            elif not cols["name"] and any(h in text for h in NAME_HINTS):
                cols["name"] = col_idx
            elif not cols["qty"] and any(h in text for h in QTY_HINTS):
                cols["qty"] = col_idx
            elif not cols["mfr"] and any(h in text for h in MFR_HINTS):
                cols["mfr"] = col_idx
        if (cols["pn"] or cols["name"]) and cols["qty"]:
            return {"row": row_idx, **cols}
        if (cols["pn"] or cols["name"]) and fallback is None:
            fallback = {"row": row_idx, **cols}
    if fallback:
        return fallback
    raise SystemExit(
        "не найдена шапка: нужна колонка с артикулом или наименованием и колонка с количеством.\n"
        "  Подсказать вручную: --header-row 3 --pn-col B --qty-col D [--mfr-col E]\n"
        "  (колонки — буквой или числом, A=1)")


def _to_qty(value) -> int | None:
    """Количество из ячейки. None — строка не позиция (примечание, подзаголовок)."""
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return int(value)
    text = str(value).strip().replace(",", ".")
    try:
        return int(float(text))
    except ValueError:
        digits = re.sub(r"[^\d]", "", text)
        return int(digits) if digits else None


def make_part(part_id: int, pn_text: str, description: str, qty: int,
              manufacturer: str = "") -> mpnmod.Part:
    """Позиция из текстов ячеек. Артикул — из pn_text (или описания), бренд — из колонки,
    иначе из текста артикула, иначе из описания."""
    source = pn_text or description
    part = mpnmod.parse_line(source, qty, part_id=part_id)
    if not part.manufacturer and description and description != source:
        part.manufacturer = mpnmod.parse_line(description, qty).manufacturer
        if part.manufacturer:
            part.brand_source = "description"
    if manufacturer:
        # колонка «Производитель» — приоритет над разбором текста (регламент)
        part.manufacturer = mpnmod.canon_brand(manufacturer)
        part.brand_source = "column"
    part.raw = description or pn_text
    return part


def read_bom(path: Path, sheet: str | None = None, header_row: int = 0, pn_col: int = 0,
             name_col: int = 0, qty_col: int = 0, mfr_col: int = 0,
             log=print) -> list[mpnmod.Part]:
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        if sheet:
            if sheet not in wb.sheetnames:
                raise SystemExit(f"в книге нет листа {sheet!r}; есть: {', '.join(wb.sheetnames)}")
            ws = wb[sheet]
        else:
            ws = wb.worksheets[0]
        found = _find_header(ws) if not (header_row and (pn_col or name_col) and qty_col) \
            else {"row": header_row, "pn": pn_col, "name": name_col, "qty": qty_col, "mfr": mfr_col, "num": 0}
    except BaseException:
        wb.close()          # на Windows незакрытая книга держит файл
        raise
    header_row = header_row or found["row"]
    pn_col = pn_col or found["pn"]
    name_col = name_col or found["name"]
    qty_col = qty_col or found["qty"]
    mfr_col = mfr_col or found["mfr"]
    log(f"Шапка в строке {header_row}. Колонки: PN={pn_col or '—'}, "
        f"Наименование={name_col or '—'}, Кол-во={qty_col or '—'}, Производитель={mfr_col or '—'}")

    parts: list[mpnmod.Part] = []
    skipped = 0
    for row in ws.iter_rows(min_row=header_row + 1):
        cells = {idx: getattr(c, "value", None) for idx, c in enumerate(row, 1)}
        pn_text = str(cells.get(pn_col) or "").strip() if pn_col else ""
        desc = str(cells.get(name_col) or "").strip() if name_col else ""
        if not pn_text and not desc:
            continue
        qty = _to_qty(cells.get(qty_col)) if qty_col else None
        if qty is None or qty <= 0:
            skipped += 1                     # примечание или подзаголовок — не позиция
            continue
        mfr = str(cells.get(mfr_col) or "").strip() if mfr_col else ""
        parts.append(make_part(len(parts) + 1, pn_text, desc, qty, mfr))
    wb.close()
    if skipped:
        log(f"Пропущено строк без количества: {skipped}")
    return parts


def read_txt(path: Path) -> list[mpnmod.Part]:
    """Резервный вход: строки «АРТИКУЛ КОЛИЧЕСТВО [ПРОИЗВОДИТЕЛЬ]», комментарии с #."""
    parts: list[mpnmod.Part] = []
    for line in path.read_text(encoding="utf-8-sig").splitlines():
        line = line.strip()
        if not line or line.startswith("#"):
            continue
        chunks = line.split()
        pn = chunks[0]
        qty = int(chunks[1]) if len(chunks) > 1 and chunks[1].isdigit() else 1
        mfr = " ".join(chunks[2:]) if len(chunks) > 2 else ""
        parts.append(literal_part(len(parts) + 1, pn, qty, mfr))
    return parts


def literal_part(part_id: int, pn: str, qty: int, manufacturer: str = "",
                 description: str = "") -> mpnmod.Part:
    """Позиция из уже чистого артикула (превью, --once, txt): разбор текста не нужен."""
    pn = mpnmod.normalize_text(pn)
    part = mpnmod.Part(part_id=part_id, raw=description or pn, qty=qty,
                       mpn=mpnmod.fix_homoglyphs(pn), mpn_raw=pn,
                       manufacturer=mpnmod.canon_brand(manufacturer),
                       brand_source="column" if manufacturer else "none")
    part.mpn_script = "cyrillic" if mpnmod.has_cyrillic(part.mpn) else "latin"
    part.variants = mpnmod.build_variants(part.mpn)
    if not part.mpn:
        part.confidence = "low"
    elif not part.manufacturer:
        part.confidence = "medium"
    return part


def autopick_request() -> Path | None:
    """Свежий Excel-запрос рядом со скриптом или в start/ (не эталон, не превью, не итог)."""
    candidates = []
    for d in (store.ROOT, store.START):
        if not d.exists():
            continue
        for f in d.glob("*.xlsx"):
            n = f.name.lower()
            if n.startswith("bom_") or n.startswith("~$") or f.name == store.PREVIEW.name:
                continue
            candidates.append(f)
    return max(candidates, key=lambda p: p.stat().st_mtime) if candidates else None


# ── превью ───────────────────────────────────────────────────────────────────

PREVIEW_COLS = ["№", "Распознанный PN (можно править)", "Кол-во", "Производитель",
                "Исходное наименование", "Уверенность", "Замечания"]
PREVIEW_WIDTHS = [6, 34, 10, 22, 50, 12, 40]


def write_preview(parts: list[mpnmod.Part], path: Path) -> Path:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Превью"
    head = PatternFill("solid", fgColor="1F4E79")
    for col, (title, width) in enumerate(zip(PREVIEW_COLS, PREVIEW_WIDTHS), 1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = head
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = width

    warn = PatternFill("solid", fgColor="FFF2CC")
    bad = PatternFill("solid", fgColor="FFC7CE")
    for row, part in enumerate(parts, 2):
        ws.cell(row=row, column=1, value=part.part_id)
        pn_cell = ws.cell(row=row, column=2, value=part.mpn or None)
        ws.cell(row=row, column=3, value=part.qty)
        ws.cell(row=row, column=4, value=part.manufacturer or None)
        ws.cell(row=row, column=5, value=part.raw or None)
        ws.cell(row=row, column=6, value=part.confidence)
        ws.cell(row=row, column=7, value="; ".join(part.warnings) or None)
        if not part.mpn:
            pn_cell.fill = bad
        elif part.confidence == "low":
            pn_cell.fill = bad
        elif not part.manufacturer:
            ws.cell(row=row, column=4).fill = warn
    ws.freeze_panes = "A2"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def read_preview(path: Path) -> list[mpnmod.Part]:
    """Превью после правок человека: PN и производитель берутся как есть, без разбора."""
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    parts: list[mpnmod.Part] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        num, pn, qty, mfr, desc = (list(row) + [None] * 5)[:5]
        if not pn and not desc:
            continue
        qty_v = _to_qty(qty)
        part = literal_part(len(parts) + 1, str(pn or "").strip(), qty_v or 1,
                            str(mfr or "").strip(), str(desc or "").strip())
        part.warnings = ["из превью"]
        parts.append(part)
    wb.close()
    return parts
