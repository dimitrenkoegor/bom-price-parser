"""
Браузерный каскад по BOM-инструментам агрегаторов (oemsecrets, Octopart,
Findchips, TrustedParts).

Три подкоманды:

    python browser_bom.py export --site oemsecrets --result "final/итог.xlsx"
        RFQ-остаток итогового файла -> CSV для загрузки в BOM Tool площадки.

    python browser_bom.py merge --site oemsecrets --export "скачанный.xlsx" \
        --result "final/итог.xlsx" [--dry-run] [--trust-echo] [--currency EUR]
        Валидирует экспорт площадки по регламенту и вливает ТОЛЬКО в RFQ-строки.

    python browser_bom.py status --result "final/итог.xlsx"
        Сводка: найдено / LLM / браузер / остаток RFQ.

Каскад не хранит состояния: остаток после площадки N — это по определению
строки, которые всё ещё RFQ в итоговом xlsx. Журнал final/browser/cascade.jsonl
пишется append-only и никогда не читается для управления потоком.

Сама браузерная нога (загрузка CSV, настройка дистрибьюторов, экспорт)
выполняется Claude через расширение Chrome — см. скилл
.claude/skills/bom-browser-cascade/SKILL.md.
"""
import argparse
import csv
import datetime
import json
import math
import os
import re
import shutil
import sys
from pathlib import Path

import openpyxl

import price_parser as parser

SCRIPT_DIR = Path(__file__).resolve().parent
BROWSER_DIR = SCRIPT_DIR / "final" / "browser"
JOURNAL_PATH = BROWSER_DIR / "cascade.jsonl"
DEFAULT_RESULT = SCRIPT_DIR / "final" / "BOM_Приложение_1.xlsx"

# ─────────────────────────────────────────────
# ПЛОЩАДКИ
# ─────────────────────────────────────────────
# У oemsecrets формат экспорта верифицирован (bomtool_to_final.py). Остальные —
# заглушки: общая карта синонимов, скорее всего, разберёт их шапку сама, а если
# нет — BrowserExportFormatError подскажет, какую строку добавить в aliases.
# currency_default=None означает «валюта неизвестна» — файл без валюты в шапке
# или колонке будет отклонён, пока оператор не передаст --currency явно.

SITES = {
    "oemsecrets": {
        "title": "oemsecrets BOM Tool",
        "url": "https://bomtool.oemsecrets.com/",
        "currency_default": "USD",  # экспорт содержит колонку 'Unit Price in USD'
        "aliases": {},
        "upload_columns": ("Part Number", "Quantity", "Manufacturer"),
    },
    "octopart": {
        "title": "Octopart BOM Tool",
        "url": "https://octopart.com/bom-tool/",
        "currency_default": None,
        "aliases": {},
        "upload_columns": ("MPN", "Qty", "Manufacturer"),
    },
    "findchips": {
        "title": "Findchips BOM",
        "url": "https://www.findchips.com/bom",
        "currency_default": None,
        "aliases": {},
        "upload_columns": ("Part Number", "Quantity", "Manufacturer"),
    },
    "trustedparts": {
        "title": "TrustedParts BOM",
        "url": "https://www.trustedparts.com/",
        "currency_default": None,
        "aliases": {},
        "upload_columns": ("Part Number", "Quantity", "Manufacturer"),
    },
}

# ─────────────────────────────────────────────
# БЕЛЫЙ/ЧЁРНЫЙ СПИСКИ ДИСТРИБЬЮТОРОВ
# ─────────────────────────────────────────────
# Белый список — единственный источник: OEMSECRETS_AUTHORIZED из price_parser.
# Чёрный список проверяется ПЕРВЫМ: "Digikey Surplus Brokers" должен быть
# отклонён несмотря на подстроку "digikey".

BROKER_DENYLIST = (
    "surplus", "broker", "excess", "obsolete", "independent", "trading",
    "traders", "stock exchange", "hard to find", "hard-to-find", "shortage",
    "franchise-free", "second source", "aftermarket", "marketplace",
)


def _authorized_distributor(name):
    """(ok, reason). Чёрный список бьёт белый; совпадение по белому — только
    на границе токена; неизвестное имя — отказ, не принятие по умолчанию."""
    norm = re.sub(r"\s+", " ", str(name or "").strip().lower())
    if not norm:
        return False, "unknown_distributor"
    for bad in BROKER_DENYLIST:
        if bad in norm:
            return False, "broker"
    for a in parser.OEMSECRETS_AUTHORIZED:
        if re.search(r"(?:^|[^a-z0-9])" + re.escape(a), norm):
            return True, None
    return False, "unknown_distributor"


# ─────────────────────────────────────────────
# ЧТЕНИЕ ИТОГОВОГО ФАЙЛА (обратная к write_results)
# ─────────────────────────────────────────────

F_LLM_RGB = "FDE9D9"      # llm_assisted (существующий)
F_BROWSER_RGB = "E4DFEC"  # browser_assisted (новый, светло-сиреневый)

_RATE_RE = re.compile(r"Курс USD/RUB \(ЦБ \+ наценка\): ([\d.]+)")
_LEAD_RE = re.compile(r"^(\d+) weeks?$")


def _fill_rgb(cell):
    rgb = getattr(getattr(cell.fill, "fgColor", None), "rgb", None)
    return str(rgb or "")[-6:].upper()


def read_results(path):
    """Читает xlsx, записанный write_results. Возвращает (results, rate_rub).
    Курс парсится из сноски и НЕ запрашивается заново — иначе уже посчитанные
    строки при каждом слиянии тихо меняли бы рублёвую цену."""
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    results, rate_rub = [], None
    for row in range(2, ws.max_row + 1):
        a = ws.cell(row=row, column=1).value
        if isinstance(a, str) and a.startswith("Найдено:"):
            m = _RATE_RE.search(a)
            if m:
                rate_rub = float(m.group(1))
            break
        vals = [ws.cell(row=row, column=c).value for c in range(1, 11)]
        if all(v in (None, "") for v in vals):
            continue
        pn = str(vals[0]).strip() if vals[0] else ""
        qty = int(float(vals[2])) if vals[2] not in (None, "") else 1
        manufacturer = str(vals[3]).strip() if vals[3] else ""
        r = {"pn": pn, "qty": qty, "manufacturer": manufacturer, "description": pn}
        if isinstance(vals[8], str) and vals[8].strip().upper() == "RFQ":
            r["status"] = "RFQ"
        else:
            lead_text = str(vals[7]).strip() if vals[7] else ""
            m = _LEAD_RE.match(lead_text)
            r.update({
                "status": "FOUND",
                "resolved_mpn": str(vals[1]).strip() if vals[1] else "",
                "distr_pn": "",
                "distributor": str(vals[4]).strip() if vals[4] else "",
                "moq": int(float(vals[5])) if vals[5] not in (None, "") else None,
                "stock": int(float(vals[6])) if vals[6] not in (None, "") else 0,
                "lead": int(m.group(1)) if m else (1 if lead_text == "1 week" else None),
                "in_stock": lead_text == "In stock",
                "price_usd": float(vals[8]) if vals[8] not in (None, "") else None,
            })
            rgb = _fill_rgb(ws.cell(row=row, column=1))
            if rgb == F_LLM_RGB:
                r["llm_assisted"] = True
            elif rgb == F_BROWSER_RGB:
                r["browser_assisted"] = True
                r["browser_source"] = "агрегатор"  # конкретная площадка — в журнале
        results.append(r)
    return results, rate_rub


def collect_pending(source):
    """[(idx, r)] строк со status=RFQ, у которых есть и PN, и производитель.
    Позиции без производителя в каскад не идут: по регламенту они RFQ навсегда,
    их загрузка на площадку даст только строки, которые придётся отклонить.
    source — путь к xlsx либо список results в памяти."""
    if isinstance(source, (str, Path)):
        source, _ = read_results(source)
    return [(i, r) for i, r in enumerate(source)
            if r.get("status") == "RFQ" and r.get("pn") and r.get("manufacturer")]


# ─────────────────────────────────────────────
# ВЫГРУЗКА RFQ-ОСТАТКА ДЛЯ ЗАГРУЗКИ НА ПЛОЩАДКУ
# ─────────────────────────────────────────────

def export_pending(source, site, out_dir=None, stem="bom"):
    """RFQ-остаток -> CSV (UTF-8 BOM) для BOM Tool площадки. Возвращает путь
    или None, если остатка нет (каскад закончен)."""
    spec = SITES[site]
    if isinstance(source, (str, Path)):
        stem = Path(source).stem
    pending = collect_pending(source)
    if not pending:
        print(f"Остаток RFQ пуст — выгружать нечего, каскад закончен.")
        return None
    out_dir = Path(out_dir) if out_dir else BROWSER_DIR
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"{stem}__{len(pending)}_{site}__upload.csv"
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(spec["upload_columns"])
        for _, r in pending:
            w.writerow([r["pn"], r.get("qty", 1), r.get("manufacturer", "")])
    print(f"Выгружено позиций: {len(pending)} (сайт: {spec['title']})")
    print(f"Файл для загрузки: {path}")
    print(f"URL площадки: {spec['url']}")
    print(f"После скачивания экспорта: python browser_bom.py merge --site {site} "
          f"--export \"<скачанный файл>\" --dry-run")
    return path


# ─────────────────────────────────────────────
# РАЗБОР ЭКСПОРТА ПЛОЩАДКИ: реестр синонимов, не четыре парсера
# ─────────────────────────────────────────────

class BrowserExportFormatError(Exception):
    pass


# Общая карта синонимов заголовков. Ключ разрешается по ПЕРВОМУ (самому
# длинному) синониму, содержащемуся в заголовке, — поэтому
# "distributor part number" уходит в sku раньше, чем "part number" заберёт
# его в req_pn. Позиционного фолбэка нет никогда: угадывание по номеру
# колонки — это ровно то, как остаток склада становится ценой.
FIELD_ALIASES = {
    "req_pn": ("part number", "part no", "requested part", "your part",
               "mpn requested", "mpn"),
    "matched_mpn": ("manufacturer part number", "matched part", "mfr part number",
                    "matched mpn", "mfr part"),
    "manufacturer": ("manufacturer", "mfr", "brand", "supplier brand"),
    "distributor": ("distributor", "seller", "supplier", "source",
                    "authorized dist"),
    "price": ("unit price", "price", "cost each", "unit cost"),
    "currency": ("currency", "curr"),
    "break_qty": ("price break", "break qty", "qty break",
                  "quantity for single bom", "quantity"),
    "moq": ("minimum order", "moq", "min qty", "min order qty"),
    "stock": ("stock", "availability", "qty available", "in stock", "inventory"),
    "lead": ("lead time", "factory lead", "lead time on additional stock"),
    "sku": ("distributor sku", "distributor part number", "vendor part", "sku"),
}
REQUIRED_FIELDS = ("req_pn", "manufacturer", "distributor", "price")
_CURRENCY_IN_HEADER = re.compile(r"\b(usd|eur|gbp)\b|[$€£]")
_CURRENCY_SYMBOL = {"$": "USD", "€": "EUR", "£": "GBP"}


def resolve_columns(headers, site):
    """Разрешает шапку экспорта в карту {field: col_idx}. Возвращает
    (colmap, header_currency). Кидает BrowserExportFormatError с подсказкой,
    если не разрешилось обязательное поле."""
    merged = {}
    site_aliases = SITES[site].get("aliases", {})
    for field, generic in FIELD_ALIASES.items():
        merged[field] = tuple(site_aliases.get(field, ())) + tuple(generic)
    # (алиас, поле) от длинного к короткому — одна колонка получает одно поле
    pairs = sorted(((a, f) for f, aliases in merged.items() for a in aliases),
                   key=lambda p: -len(p[0]))
    colmap, header_currency = {}, None
    for idx, h in enumerate(headers):
        hl = re.sub(r"\s+", " ", str(h or "").strip().lower())
        if not hl:
            continue
        for alias, field in pairs:
            if alias in hl:
                if field not in colmap:          # левая колонка выигрывает
                    colmap[field] = idx
                    if field == "price":
                        m = _CURRENCY_IN_HEADER.search(hl)
                        if m:
                            token = m.group(0)
                            header_currency = _CURRENCY_SYMBOL.get(token, token.upper())
                break
    missing = [f for f in REQUIRED_FIELDS if f not in colmap]
    if missing:
        raise BrowserExportFormatError(
            f"Не разрешены обязательные поля {missing} для площадки '{site}'.\n"
            f"Заголовки в файле: {[str(h) for h in headers if h]}\n"
            f"Добавьте синонимы в SITES['{site}']['aliases'], например: "
            f"{{'{missing[0]}': ('<заголовок из файла>',)}}. Ничего не импортировано.")
    return colmap, header_currency


def _num(value):
    if value in (None, ""):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(" ", "").replace(" ", "")
    text = text.replace(",", ".")
    m = re.search(r"-?\d+(?:\.\d+)?", text)
    return float(m.group(0)) if m else None


def _read_table(path):
    """xlsx или csv -> список строк (списков значений)."""
    p = Path(path)
    if p.suffix.lower() in (".xlsx", ".xlsm"):
        wb = openpyxl.load_workbook(p, data_only=True)
        ws = wb.active
        return [[c for c in row] for row in ws.iter_rows(values_only=True)]
    with open(p, newline="", encoding="utf-8-sig") as f:
        return [row for row in csv.reader(f)]


def read_site_export(path, site):
    """Экспорт площадки -> (rows, info). rows — нормализованные словари
    {req_pn, matched_mpn, manufacturer, distributor, price, currency,
    break_qty, moq, stock, lead, sku}. info.echo_only=True, если в экспорте
    нет отдельной колонки подобранного MPN (валидация MPN — тавтология)."""
    table = _read_table(path)
    if not table:
        raise BrowserExportFormatError(f"Пустой файл экспорта: {path}")
    # шапка может быть не в первой строке (баннер) — ищем в первых 5 строках
    best = None
    for i, row in enumerate(table[:5]):
        try:
            colmap, cur = resolve_columns(row, site)
        except BrowserExportFormatError:
            continue
        if best is None or len(colmap) > len(best[1]):
            best = (i, colmap, cur)
    if best is None:
        resolve_columns(table[0], site)          # кинет ошибку с подсказкой
    header_row, colmap, header_currency = best
    echo_only = "matched_mpn" not in colmap

    def cell(row, field):
        idx = colmap.get(field)
        return row[idx] if idx is not None and idx < len(row) else None

    rows = []
    for raw in table[header_row + 1:]:
        req_pn = str(cell(raw, "req_pn") or "").strip()
        if not req_pn:
            continue
        cur = str(cell(raw, "currency") or "").strip().upper() or header_currency
        rows.append({
            "req_pn": req_pn,
            "matched_mpn": str(cell(raw, "matched_mpn") or "").strip(),
            "manufacturer": str(cell(raw, "manufacturer") or "").strip(),
            "distributor": str(cell(raw, "distributor") or "").strip(),
            "price": _num(cell(raw, "price")),
            "currency": cur or None,
            "break_qty": _num(cell(raw, "break_qty")),
            "moq": _num(cell(raw, "moq")),
            "stock": _num(cell(raw, "stock")),
            "lead": _num(cell(raw, "lead")),
            "sku": str(cell(raw, "sku") or "").strip(),
        })
    info = {"echo_only": echo_only, "header_currency": header_currency,
            "colmap": colmap, "header_row": header_row}
    return rows, info


# ─────────────────────────────────────────────
# ВАЛЮТА (fail-closed)
# ─────────────────────────────────────────────

def _to_usd(value, currency, rates):
    """USD — тождество; EUR — по курсу, взятому один раз на слияние.
    Прочее/неизвестное -> None (bad_currency), никакого тихого угадывания."""
    if currency == "USD":
        return value
    if currency == "EUR":
        if "EUR" not in rates:
            rates["EUR"] = parser.get_eur_usd()
        return value * rates["EUR"]
    return None


# ─────────────────────────────────────────────
# ШЛЮЗ ВАЛИДАЦИИ (порядок важен, первый отказ выигрывает)
# ─────────────────────────────────────────────

def validate_row(req, row, *, rates, default_currency=None, forced_currency=None,
                 max_unit_price=10000.0):
    """(offer|None, reason). offer — стандартный контракт дистрибьютора
    (price_parser.py:256), дальше выбор делает parser._best_candidate."""
    # 1. Производитель обязателен. Строже, чем oemsecrets_pass (тот пропускает
    #    пустого производителя, потому что API-результат сам по себе доверенный);
    #    у браузерных данных подтверждения API нет.
    if not row.get("manufacturer"):
        return None, "no_manufacturer"
    # 2. Чужой производитель = аналог, запрещено регламентом.
    if not parser.manufacturer_matches(req.get("manufacturer", ""), row["manufacturer"]):
        return None, "wrong_manufacturer"
    # 3. MPN: маска x, суффиксы упаковки — через существующий mpn_matches.
    cand_mpn = row.get("matched_mpn") or row["req_pn"]
    if not parser.mpn_matches(req["pn"], cand_mpn):
        return None, "analog_mpn"
    # 4. Дистрибьютор: чёрный список -> белый список -> отказ по умолчанию.
    ok, reason = _authorized_distributor(row.get("distributor"))
    if not ok:
        return None, reason
    # 5. Цена: конечная, положительная, не абсурдная.
    price = row.get("price")
    if price is None or not math.isfinite(price) or price <= 0 or price > max_unit_price:
        return None, "bad_price"
    # 6. Валюта: строка > заголовок > дефолт площадки; дефолт не-USD без
    #    явного --currency -> отказ.
    cur = row.get("currency") or forced_currency or default_currency
    if not cur or (cur != "USD" and row.get("currency") is None
                   and forced_currency is None and default_currency != "USD"):
        return None, "bad_currency"
    usd = _to_usd(price, cur, rates)
    if usd is None:
        return None, "bad_currency"
    break_qty = int(row["break_qty"]) if row.get("break_qty") else 1
    offer = {
        "dist": row["distributor"],
        "pn": row.get("sku") or cand_mpn,
        "mpn": cand_mpn,
        "stock": int(row["stock"]) if row.get("stock") else 0,
        "moq": int(row["moq"]) if row.get("moq") else break_qty,
        "lead": int(row["lead"]) if row.get("lead") else None,
        "tiers": [{"qty": break_qty, "price": round(usd, 5)}],
        "manu": row["manufacturer"],
        "_currency": cur,
    }
    return offer, "accepted"


# ─────────────────────────────────────────────
# СЛИЯНИЕ (структурный близнец oemsecrets_pass)
# ─────────────────────────────────────────────

def _journal(journal_path, record):
    if not journal_path:
        return
    Path(journal_path).parent.mkdir(parents=True, exist_ok=True)
    with open(journal_path, "a", encoding="utf-8") as f:
        f.write(json.dumps(record, ensure_ascii=False) + "\n")


def browser_merge_pass(results, rows, site, *, trust_echo=False, echo_only=False,
                       default_currency=None, forced_currency=None,
                       max_unit_price=10000.0, rates=None,
                       journal_path=None, verbose=True):
    """Идёт ТОЛЬКО по RFQ-остатку, валидирует строки экспорта, пишет принятое
    в results[i] с флагом browser_assisted. Возвращает stats."""
    rates = rates if rates is not None else {}
    ts = datetime.datetime.now().isoformat(timespec="seconds")
    stats = {"accepted": 0, "not_pending": 0, "rejected": {}}

    def reject(reason, row):
        stats["rejected"][reason] = stats["rejected"].get(reason, 0) + 1
        _journal(journal_path, {"ts": ts, "site": site, "req_pn": row.get("req_pn"),
                                "matched_mpn": row.get("matched_mpn"),
                                "distributor": row.get("distributor"),
                                "price": row.get("price"), "currency": row.get("currency"),
                                "verdict": "rejected", "reason": reason})

    if echo_only and not trust_echo:
        # В экспорте нет отдельной колонки подобранного MPN: проверка MPN была
        # бы тавтологией, любой аналог прошёл бы. Отказ файлу целиком.
        print(f"    ОТКАЗ ФАЙЛУ: экспорт '{site}' не содержит колонки подобранного "
              f"MPN (echo-only). Повторите с --trust-echo, если производитель в "
              f"экспорте заполнен и вы принимаете сопоставление площадки.")
        for row in rows:
            reject("echo_only", row)
        return stats

    pending = {parser.canonical_mpn_loose(r["pn"]): (i, r)
               for i, r in collect_pending(results)}
    grouped = {}
    for row in rows:
        grouped.setdefault(parser.canonical_mpn_loose(row["req_pn"]), []).append(row)

    converted = 0
    for key, group in grouped.items():
        if key not in pending:
            stats["not_pending"] += len(group)
            for row in group:
                _journal(journal_path, {"ts": ts, "site": site,
                                        "req_pn": row.get("req_pn"),
                                        "verdict": "rejected", "reason": "not_pending"})
            continue
        i, req = pending[key]
        offers = []
        for row in group:
            offer, reason = validate_row(
                req, row, rates=rates, default_currency=default_currency,
                forced_currency=forced_currency, max_unit_price=max_unit_price)
            if offer:
                offers.append(offer)
                if offer["_currency"] != "USD":
                    converted += 1
            else:
                reject(reason, row)
        if not offers:
            continue
        # Регламент §4: при равной цене — больший сток. Ключ _best_candidate
        # только (in_stock, price) и сравнение строгое, поэтому достаточно
        # предварительной сортировки по убыванию стока.
        offers.sort(key=lambda o: -o["stock"])
        cand = parser._best_candidate(offers, req["pn"], req.get("qty", 1))
        cur = next(o["_currency"] for o in offers
                   if o["dist"] == cand["distributor"])
        assert results[i].get("status") == "RFQ", \
            f"merge пытается перезаписать не-RFQ строку #{i} ({req['pn']})"
        cand["description"] = req.get("description", req["pn"])
        cand["manufacturer"] = cand.get("manufacturer") or req.get("manufacturer", "")
        cand["browser_assisted"] = True
        cand["browser_source"] = site
        note = f"{site}, {cur}"
        if cur == "EUR":
            note += f"->USD @{rates['EUR']:.4f}"
        cand["browser_note"] = note[:120]
        results[i] = cand
        stats["accepted"] += 1
        _journal(journal_path, {"ts": ts, "site": site, "req_pn": req["pn"],
                                "matched_mpn": cand.get("resolved_mpn"),
                                "distributor": cand["distributor"],
                                "price": cand["price_usd"], "currency": cur,
                                "verdict": "accepted", "reason": "accepted"})
        if verbose:
            print(f"    BOM+{site}: {req['pn']} -> {cand.get('resolved_mpn') or req['pn']} "
                  f"| {cand['distributor']} ${cand['price_usd']}")
    if converted and "EUR" in rates:
        print(f"    ВНИМАНИЕ: {converted} строк конвертировано EUR->USD "
              f"по курсу {rates['EUR']:.4f}")
    rej = sum(stats["rejected"].values())
    print(f"Браузерный контур ({site}): принято {stats['accepted']}, "
          f"отклонено {rej}, вне остатка {stats['not_pending']}")
    if stats["rejected"]:
        detail = ", ".join(f"{k}: {v}" for k, v in sorted(stats["rejected"].items()))
        print(f"    Причины отказов: {detail}")
    return stats


def merge_export(result_path, export_path, site, *, dry_run=False, trust_echo=False,
                 forced_currency=None, max_unit_price=10000.0):
    """read_results -> browser_merge_pass -> write_results через временный файл
    с атомарной заменой и .bak-копией. --dry-run ничего не пишет."""
    result_path = Path(result_path)
    results, rate_rub = read_results(result_path)
    before = sum(1 for r in results if r.get("status") == "RFQ")
    rows, info = read_site_export(export_path, site)
    print(f"Экспорт {SITES[site]['title']}: строк {len(rows)}, "
          f"echo-only: {'да' if info['echo_only'] else 'нет'}, "
          f"валюта шапки: {info['header_currency'] or '—'}")
    stats = browser_merge_pass(
        results, rows, site,
        trust_echo=trust_echo, echo_only=info["echo_only"],
        default_currency=SITES[site]["currency_default"],
        forced_currency=forced_currency, max_unit_price=max_unit_price,
        journal_path=None if dry_run else JOURNAL_PATH)
    after = sum(1 for r in results if r.get("status") == "RFQ")
    print(f"RFQ: было {before} -> станет {after}")
    if dry_run:
        print("Режим --dry-run: файл не изменён.")
        return stats
    if stats["accepted"] == 0:
        print("Принятых строк нет — файл не изменён.")
        return stats
    if rate_rub is None:
        print("Курс не найден в сноске — запрашиваю заново (ЦБ + наценка 4).")
        rate_rub = parser.get_usd_rub()
    bak = result_path.with_name(result_path.stem + ".bak.xlsx")
    shutil.copy2(result_path, bak)
    tmp = result_path.with_name(result_path.stem + ".tmp.xlsx")
    parser.write_results(results, tmp, rate_rub)
    os.replace(tmp, result_path)
    print(f"Резервная копия: {bak}")
    print(f"Итог обновлён: {result_path}")
    return stats


def print_status(result_path):
    results, rate_rub = read_results(result_path)
    found = sum(1 for r in results if r.get("status") == "FOUND")
    llm = sum(1 for r in results if r.get("llm_assisted"))
    br = sum(1 for r in results if r.get("browser_assisted"))
    pend = collect_pending(results)
    no_mfr = sum(1 for r in results
                 if r.get("status") == "RFQ" and not r.get("manufacturer"))
    print(f"Всего позиций: {len(results)}")
    print(f"Найдено: {found} (API: {found - llm - br}, LLM: {llm}, браузер: {br})")
    print(f"RFQ: {len(results) - found}, из них в каскад годно: {len(pend)} "
          f"(без производителя: {no_mfr})")
    print(f"Курс из сноски: {rate_rub if rate_rub is not None else 'не найден'}")


# ─────────────────────────────────────────────
# CLI
# ─────────────────────────────────────────────

def main():
    ap = argparse.ArgumentParser(
        description="Браузерный каскад по BOM-инструментам агрегаторов")
    sub = ap.add_subparsers(dest="cmd", required=True)

    p_exp = sub.add_parser("export", help="RFQ-остаток -> CSV для загрузки")
    p_exp.add_argument("--site", required=True, choices=sorted(SITES))
    p_exp.add_argument("--result", default=str(DEFAULT_RESULT))
    p_exp.add_argument("--out-dir", default=None)

    p_mrg = sub.add_parser("merge", help="Влить экспорт площадки в итоговый xlsx")
    p_mrg.add_argument("--site", required=True, choices=sorted(SITES))
    p_mrg.add_argument("--result", default=str(DEFAULT_RESULT))
    p_mrg.add_argument("--export", required=True, help="Скачанный экспорт площадки")
    p_mrg.add_argument("--dry-run", action="store_true",
                       help="Показать принято/отклонено, ничего не писать")
    p_mrg.add_argument("--trust-echo", action="store_true",
                       help="Принять экспорт без колонки подобранного MPN")
    p_mrg.add_argument("--currency", default=None, choices=("USD", "EUR"),
                       help="Валюта строк без явной валюты в файле")
    p_mrg.add_argument("--max-unit-price", type=float, default=10000.0)

    p_st = sub.add_parser("status", help="Сводка по итоговому файлу")
    p_st.add_argument("--result", default=str(DEFAULT_RESULT))

    args = ap.parse_args()
    if args.cmd == "export":
        export_pending(args.result, args.site, out_dir=args.out_dir)
    elif args.cmd == "merge":
        merge_export(args.result, args.export, args.site,
                     dry_run=args.dry_run, trust_echo=args.trust_echo,
                     forced_currency=args.currency,
                     max_unit_price=args.max_unit_price)
    else:
        print_status(args.result)


if __name__ == "__main__":
    main()
